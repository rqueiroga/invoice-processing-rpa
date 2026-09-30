from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parent
QUEUE = ROOT / "data" / "invoice_queue.xlsx"
STATE = ROOT / "runtime" / "state.json"
UPLOADS = ROOT / "runtime" / "uploads"


def reset() -> None:
    workbook = load_workbook(QUEUE)
    sheet = workbook["Invoices"]
    headers = {sheet.cell(1, col).value: col for col in range(1, sheet.max_column + 1)}

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, headers["Automation Status"]).value = ""

    workbook.save(QUEUE)
    workbook.close()

    STATE.parent.mkdir(parents=True, exist_ok=True)
    STATE.write_text(json.dumps({"registrations": {}}, indent=2), encoding="utf-8")
    shutil.rmtree(UPLOADS, ignore_errors=True)


if __name__ == "__main__":
    reset()
    print("Demo reset complete.")
