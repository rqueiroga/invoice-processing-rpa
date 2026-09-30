from __future__ import annotations

import argparse
import logging
import os
import shutil
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).parent
DEFAULT_FILE = ROOT / "data" / "invoice_queue.xlsx"
SUPPLIER_URL = os.getenv("SUPPLIER_URL", "http://127.0.0.1:8000/supplier")
FINANCE_URL = os.getenv("FINANCE_URL", "http://127.0.0.1:8000/finance")
SUPPLIER_USER = os.getenv("SUPPLIER_USER", "supplier.demo")
SUPPLIER_PASSWORD = os.getenv("SUPPLIER_PASSWORD", "demo123")
FINANCE_USER = os.getenv("FINANCE_USER", "finance.demo")
FINANCE_PASSWORD = os.getenv("FINANCE_PASSWORD", "demo123")
DEMO_SPEED_MS = int(os.getenv("DEMO_SPEED_MS", "600"))
DEMO_LANG = os.getenv("DEMO_LANG", "en").lower()

TYPE_MAP = {
    "Service Invoice": "Services",
    "Product Invoice": "Goods",
    "Expense Reimbursement": "Expenses",
}

DISPLAY_TYPES_PT = {
    "Service Invoice": "Nota de Serviço",
    "Product Invoice": "Nota de Produto",
    "Expense Reimbursement": "Reembolso de Despesas",
    "Other": "Outro",
}


def display_status_pt(status: str) -> str:
    translations = {
        "Completed": "Concluído",
        "Purchase order not found": "Pedido de compra não encontrado",
        "Already registered": "Já cadastrado",
        "Documents not found": "Documentos não encontrados",
    }
    if status.startswith("Unmapped invoice type:"):
        invoice_type = status.split(":", 1)[1].strip()
        return f"Tipo de nota sem mapeamento: {DISPLAY_TYPES_PT.get(invoice_type, invoice_type)}"
    if status.startswith("Error:"):
        return "Erro:" + status.split(":", 1)[1]
    return translations.get(status, status)

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(message)s")
LOG = logging.getLogger("invoice_rpa")


@dataclass(frozen=True)
class Invoice:
    row: int
    invoice_id: str
    purchase_order: str
    due_date: date
    invoice_type: str


class InvoiceSheet:
    HEADERS = {
        "invoice_id": "Invoice ID",
        "purchase_order": "Purchase Order",
        "due_date": "Due Date",
        "invoice_type": "Invoice Type",
        "status": "Automation Status",
    }

    def __init__(self, path: Path):
        self.path = path
        self.workbook = load_workbook(path)
        self.sheet = self.workbook["Invoices"]
        names = {self.sheet.cell(1, col).value: col for col in range(1, self.sheet.max_column + 1)}
        self.cols = {key: names[label] for key, label in self.HEADERS.items()}

    def rows(self):
        for row in range(2, self.sheet.max_row + 1):
            if self.sheet.cell(row, self.cols["invoice_id"]).value:
                yield row

    def read(self, row: int) -> Invoice:
        value = self.sheet.cell(row, self.cols["due_date"]).value
        due = value.date() if isinstance(value, datetime) else value
        return Invoice(
            row=row,
            invoice_id=str(self.sheet.cell(row, self.cols["invoice_id"]).value).strip(),
            purchase_order=str(self.sheet.cell(row, self.cols["purchase_order"]).value).strip(),
            due_date=due,
            invoice_type=str(self.sheet.cell(row, self.cols["invoice_type"]).value).strip(),
        )

    def status(self, row: int) -> str:
        return str(self.sheet.cell(row, self.cols["status"]).value or "").strip()

    def save_status(self, row: int, status: str) -> None:
        self.sheet.cell(row, self.cols["status"]).value = status
        self.workbook.save(self.path)

    def close(self) -> None:
        self.workbook.close()


class SupplierPortal:
    def __init__(self, page):
        self.page = page

    def login(self) -> None:
        LOG.info("Portal do Fornecedor | fazendo login")
        self.page.bring_to_front()
        self.page.goto(f"{SUPPLIER_URL}/login")
        self.page.get_by_test_id("user").fill(SUPPLIER_USER)
        self.page.get_by_test_id("password").fill(SUPPLIER_PASSWORD)
        self.page.get_by_test_id("sign-in").click()
        self.page.wait_for_url("**/supplier/invoices")

    def download_package(self, invoice_id: str, folder: Path) -> Path | None:
        LOG.info("Portal do Fornecedor | buscando %s", invoice_id)
        self.page.bring_to_front()
        self.page.goto(f"{SUPPLIER_URL}/invoices?invoice={quote(invoice_id)}")

        link = self.page.get_by_test_id("invoice-link")
        if link.count() == 0:
            LOG.info("Portal do Fornecedor | nota não encontrada")
            return None

        link.click()
        rows = self.page.get_by_test_id("attachment-row")
        if rows.count() < 2:
            LOG.info("Portal do Fornecedor | documentos insuficientes")
            return None

        LOG.info("Portal do Fornecedor | baixando o primeiro e o último documento")
        files = []
        for label, index in (("first", 0), ("last", rows.count() - 1)):
            row = rows.nth(index)
            with self.page.expect_download() as event:
                row.get_by_test_id("download").click()
            download = event.value
            target = folder / f"{label}_{download.suggested_filename}"
            download.save_as(target)
            files.append(target)

        output = folder / f"{invoice_id}_package.pdf"
        LOG.info("PDF | unificando documentos em %s", output.name)
        return merge_pdfs(files, output)


class FinancePortal:
    def __init__(self, page):
        self.page = page

    def login(self) -> None:
        LOG.info("Portal Financeiro | fazendo login")
        self.page.bring_to_front()
        self.page.goto(f"{FINANCE_URL}/login")
        self.page.get_by_test_id("user").fill(FINANCE_USER)
        self.page.get_by_test_id("password").fill(FINANCE_PASSWORD)
        self.page.get_by_test_id("sign-in").click()
        self.page.wait_for_url("**/finance/orders")

    def register(self, invoice: Invoice, document: Path) -> str:
        doc_type = TYPE_MAP.get(invoice.invoice_type)
        if not doc_type:
            LOG.info("Regra de negócio | tipo de nota sem mapeamento: %s", invoice.invoice_type)
            return f"Unmapped invoice type: {invoice.invoice_type}"

        LOG.info("Portal Financeiro | buscando pedido de compra %s", invoice.purchase_order)
        self.page.bring_to_front()
        self.page.goto(f"{FINANCE_URL}/orders?po={quote(invoice.purchase_order)}")

        order = self.page.get_by_test_id("order-link")
        if order.count() == 0:
            LOG.info("Portal Financeiro | pedido de compra não encontrado")
            return "Purchase order not found"

        order.click()
        LOG.info("Portal Financeiro | verificando registro duplicado")
        if self.page.locator(f"[data-invoice='{invoice.invoice_id}']").count():
            return "Already registered"

        LOG.info("Portal Financeiro | cadastrando pacote da nota")
        self.page.get_by_test_id("register-link").click()
        self.page.get_by_test_id("invoice-id").fill(invoice.invoice_id)
        self.page.get_by_test_id("document-type").select_option(value=doc_type)
        self.page.get_by_test_id("processing-date").fill(date.today().isoformat())
        self.page.get_by_test_id("pdf-package").set_input_files(document)
        self.page.get_by_test_id("save-registration").click()
        self.page.get_by_test_id("success-message").wait_for()
        return "Completed"


def merge_pdfs(files: list[Path], output: Path) -> Path:
    writer = PdfWriter()
    for file in files:
        for page in PdfReader(file).pages:
            writer.add_page(page)
    with output.open("wb") as handle:
        writer.write(handle)
    return output


def browser_for(playwright, headless: bool):
    executable = os.getenv("BROWSER_PATH") or shutil.which("chromium") or shutil.which("google-chrome")
    options = {"headless": headless, "slow_mo": DEMO_SPEED_MS}
    if executable:
        options["executable_path"] = executable
    return playwright.chromium.launch(**options)


def run(path: Path, reprocess: bool, headless: bool) -> None:
    sheet = InvoiceSheet(path)
    try:
        with sync_playwright() as playwright, tempfile.TemporaryDirectory(prefix="invoice_rpa_") as temp:
            browser = browser_for(playwright, headless)
            context = browser.new_context(accept_downloads=True)

            if DEMO_LANG == "pt":
                context.add_cookies([{"name": "lang", "value": "pt", "url": "http://127.0.0.1:8000"}])

            supplier = SupplierPortal(context.new_page())
            finance = FinancePortal(context.new_page())
            supplier.login()
            finance.login()

            for row in sheet.rows():
                if sheet.status(row) and not reprocess:
                    continue

                invoice = sheet.read(row)
                LOG.info(
                    "Processando | %s | %s | %s",
                    invoice.invoice_id,
                    invoice.purchase_order,
                    DISPLAY_TYPES_PT.get(invoice.invoice_type, invoice.invoice_type),
                )

                try:
                    folder = Path(temp) / invoice.invoice_id
                    folder.mkdir(parents=True, exist_ok=True)
                    package = supplier.download_package(invoice.invoice_id, folder)
                    result = "Documents not found" if package is None else finance.register(invoice, package)
                except Exception as exc:
                    LOG.exception("Falha na automação para %s", invoice.invoice_id)
                    result = f"Error: {str(exc)[:180]}"

                sheet.save_status(row, result)
                LOG.info("Resultado | %s -> %s", invoice.invoice_id, display_status_pt(result))

            if not headless:
                LOG.info("Demonstração concluída | mantendo o navegador aberto por 5 segundos")
                finance.page.bring_to_front()
                finance.page.wait_for_timeout(5000)

            context.close()
            browser.close()
    finally:
        sheet.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Demonstração de RPA para processamento de notas")
    parser.add_argument("--file", type=Path, default=DEFAULT_FILE)
    parser.add_argument("--reprocess", action="store_true")
    parser.add_argument("--headless", action="store_true")
    return parser.parse_args()


if __name__ == "__main__":
    options = parse_args()
    run(options.file, options.reprocess, options.headless)
